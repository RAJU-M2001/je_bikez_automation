"""
JE BikeZ - Shared Excel and HTML reporting.
"""

import html
import logging
import time
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from web_automation.config.settings import ensure_report_dir

LOGGER_NAME = "web_automation"
logger = logging.getLogger(LOGGER_NAME)

TEST_CASE_RESULTS = []


def get_logger(name: str = LOGGER_NAME):
    """Return a configured logger."""
    return logging.getLogger(name)


def reset_test_case_results() -> None:
    """Clear individual test-case results before a new run."""
    TEST_CASE_RESULTS.clear()


def log_test_case(
    test_name: str,
    test_case: str,
    status: str,
    message: str = "",
    **kwargs,
) -> None:
    """Record one individual validation for the reports."""
    normalized_status = (
        "PASSED" if status.upper() in {"PASS", "PASSED"} else "FAILED"
    )

    TEST_CASE_RESULTS.append(
        {
            "test_name": test_name,
            "test_case": test_case,
            "status": normalized_status,
            "message": message,
            "duration": kwargs.get("duration", ""),
        }
    )


def _report_directory() -> Path:
    """Resolve the report directory from project settings."""
    directory = ensure_report_dir()
    return Path(directory)


def _normalize_test_results(results):
    """Normalize runner-level results into dictionaries."""
    normalized = []

    for item in results or []:
        if isinstance(item, dict):
            normalized.append(item)
        else:
            normalized.append(
                {
                    "test_name": str(item),
                    "status": "UNKNOWN",
                    "message": "",
                }
            )

    return normalized


def save_test_report(results, total_duration: float = 0.0) -> str:
    """
    Create an Excel workbook containing Summary, Test Results,
    and Test Cases worksheets.
    """
    report_dir = _report_directory()
    report_dir.mkdir(parents=True, exist_ok=True)

    output_path = report_dir / "automation_report.xlsx"
    runner_results = _normalize_test_results(results)

    workbook = Workbook()

    summary = workbook.active
    summary.title = "Summary"
    summary.append(["JE BikeZ Automation Report"])
    summary.append(["Total Duration (seconds)", round(total_duration, 2)])
    summary.append(["Individual Test Cases", len(TEST_CASE_RESULTS)])
    summary.append(
        [
            "Passed Test Cases",
            sum(r["status"] == "PASSED" for r in TEST_CASE_RESULTS),
        ]
    )
    summary.append(
        [
            "Failed Test Cases",
            sum(r["status"] == "FAILED" for r in TEST_CASE_RESULTS),
        ]
    )

    test_results = workbook.create_sheet("Test Results")
    test_results.append(["Test Name", "Status", "Message"])

    for result in runner_results:
        test_results.append(
            [
                result.get("test_name", result.get("name", "")),
                result.get("status", ""),
                result.get("message", result.get("error", "")),
            ]
        )

    test_cases = workbook.create_sheet("Test Cases")
    test_cases.append(
        ["Test Name", "Test Case", "Status", "Message", "Duration"]
    )

    for result in TEST_CASE_RESULTS:
        test_cases.append(
            [
                result["test_name"],
                result["test_case"],
                result["status"],
                result["message"],
                result["duration"],
            ]
        )

    # Basic report formatting.
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"

        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(
                fill_type="solid",
                fgColor="243B53",
            )

        for column_cells in sheet.columns:
            max_length = max(
                len(str(cell.value or ""))
                for cell in column_cells
            )
            column_letter = get_column_letter(column_cells[0].column)
            sheet.column_dimensions[column_letter].width = min(
                max(max_length + 2, 14),
                70,
            )

    workbook.save(output_path)
    logger.info("Excel report saved: %s", output_path)

    return str(output_path)


def save_html_report(
    results,
    total_duration: float = 0.0,
    test_case_results=None,
) -> str:
    """
    Generate the HTML test report using the previous dashboard-style UI.

    Individual validation cases are intentionally excluded from HTML.
    They remain available in the Excel Test Cases worksheet.
    """
    from datetime import datetime
    from html import escape
    from pathlib import Path

    report_dir = Path(ensure_report_dir())
    report_dir.mkdir(parents=True, exist_ok=True)

    output_path = report_dir / "pytest_Report.html"

    # Normalize runner-level test results.
    normalized_results = []

    for result in results or []:
        if isinstance(result, dict):
            name = result.get(
                "test_name",
                result.get("name", "Unnamed Test"),
            )
            raw_status = str(result.get("status", "UNKNOWN")).upper()
            message = result.get(
                "message",
                result.get("error", ""),
            )
            duration = result.get("duration", 0)
        else:
            name = str(result)
            raw_status = "UNKNOWN"
            message = ""
            duration = 0

        if raw_status in ("PASS", "PASSED", "SUCCESS", "SUCCESSFUL"):
            status = "PASSED"
        elif raw_status in ("FAIL", "FAILED", "ERROR"):
            status = "FAILED"
        else:
            status = raw_status

        normalized_results.append(
            {
                "name": name,
                "status": status,
                "message": message,
                "duration": duration,
            }
        )

    total_tests = len(normalized_results)
    passed_tests = sum(
        item["status"] == "PASSED"
        for item in normalized_results
    )
    failed_tests = sum(
        item["status"] == "FAILED"
        for item in normalized_results
    )
    other_tests = total_tests - passed_tests - failed_tests

    if failed_tests:
        overall_status = "FAILED"
    elif total_tests and passed_tests == total_tests:
        overall_status = "PASSED"
    else:
        overall_status = "NO RESULTS"

    generated_at = datetime.now().astimezone().strftime(
        "%d %b %Y, %I:%M:%S %p %Z"
    )

    def safe(value):
        return escape(str(value if value is not None else ""))

    def format_duration(value):
        try:
            return f"{float(value):.2f}s"
        except (TypeError, ValueError):
            return "—"

    # Only top-level test results appear in this HTML table.
    result_rows = ""

    for item in normalized_results:
        status_class = (
            "passed"
            if item["status"] == "PASSED"
            else "failed"
            if item["status"] == "FAILED"
            else "other"
        )

        result_rows += f"""
        <tr>
            <td class="test-name">{safe(item["name"])}</td>
            <td>
                <span class="status {status_class}">
                    {safe(item["status"])}
                </span>
            </td>
            <td>{safe(format_duration(item["duration"]))}</td>
            <td class="message">{safe(item["message"])}</td>
        </tr>
        """

    if not result_rows:
        result_rows = """
        <tr>
            <td colspan="4" class="empty">
                No test results are available for this run.
            </td>
        </tr>
        """

    overall_class = (
        "passed"
        if overall_status == "PASSED"
        else "failed"
        if overall_status == "FAILED"
        else "other"
    )

    html_document = f"""<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1">

    <title>JE BikeZ | Automation Test Report</title>

    <style>
        :root {{
            --background: #f4f6fa;
            --surface: #ffffff;
            --text: #172033;
            --muted: #697586;
            --border: #e3e8ef;
            --accent: #ff6600;
            --green: #16845b;
            --red: #d64545;
            --blue: #315b96;
        }}

        * {{
            box-sizing: border-box;
        }}

        body {{
            margin: 0;
            padding: 32px;
            background: var(--background);
            color: var(--text);
            font-family: Inter, "Segoe UI", Arial, sans-serif;
        }}

        .container {{
            max-width: 1240px;
            margin: 0 auto;
        }}

        .header {{
            display: flex;
            justify-content: space-between;
            align-items: flex-start;
            gap: 20px;
            margin-bottom: 28px;
        }}

        .eyebrow {{
            color: var(--accent);
            font-size: 12px;
            font-weight: 800;
            letter-spacing: 1.6px;
            text-transform: uppercase;
        }}

        h1 {{
            margin: 8px 0;
            font-size: 30px;
            letter-spacing: -0.8px;
        }}

        .subtitle {{
            margin: 0;
            color: var(--muted);
            font-size: 14px;
        }}

        .overall {{
            display: inline-block;
            padding: 9px 14px;
            border-radius: 8px;
            font-size: 12px;
            font-weight: 800;
            letter-spacing: 0.5px;
        }}

        .overall.passed {{
            color: var(--green);
            background: #e6f6ee;
        }}

        .overall.failed {{
            color: var(--red);
            background: #fdeaea;
        }}

        .overall.other {{
            color: var(--blue);
            background: #eaf1fb;
        }}

        .summary {{
            display: grid;
            grid-template-columns: repeat(5, minmax(0, 1fr));
            gap: 16px;
            margin-bottom: 28px;
        }}

        .card {{
            padding: 21px;
            background: var(--surface);
            border: 1px solid var(--border);
            border-radius: 12px;
            box-shadow: 0 3px 12px rgba(20, 35, 60, 0.035);
        }}

        .card-label {{
            color: var(--muted);
            font-size: 12px;
            font-weight: 600;
        }}

        .card-value {{
            margin-top: 12px;
            font-size: 27px;
            font-weight: 750;
            overflow-wrap: anywhere;
        }}

        .card-value.green {{
            color: var(--green);
        }}

        .card-value.red {{
            color: var(--red);
        }}

        .panel {{
            overflow: hidden;
            background: var(--surface);
            border: 1px solid var(--border);
            border-radius: 12px;
            box-shadow: 0 3px 12px rgba(20, 35, 60, 0.035);
        }}

        .panel-header {{
            display: flex;
            justify-content: space-between;
            align-items: center;
            gap: 12px;
            padding: 22px 24px;
            border-bottom: 1px solid var(--border);
        }}

        h2 {{
            margin: 0;
            font-size: 18px;
        }}

        .panel-subtitle {{
            margin-top: 6px;
            color: var(--muted);
            font-size: 13px;
        }}

        .table-wrap {{
            overflow-x: auto;
        }}

        table {{
            width: 100%;
            border-collapse: collapse;
            text-align: left;
        }}

        thead {{
            background: #f8f9fc;
        }}

        th {{
            padding: 14px 20px;
            color: #586477;
            font-size: 11px;
            font-weight: 750;
            letter-spacing: 0.7px;
            text-transform: uppercase;
            white-space: nowrap;
        }}

        td {{
            padding: 17px 20px;
            border-top: 1px solid #edf0f5;
            font-size: 13px;
            vertical-align: top;
        }}

        tbody tr:hover {{
            background: #fafbfe;
        }}

        .test-name {{
            min-width: 150px;
            font-weight: 650;
        }}

        .message {{
            min-width: 220px;
            color: var(--muted);
            overflow-wrap: anywhere;
        }}

        .status {{
            display: inline-block;
            padding: 6px 9px;
            border-radius: 6px;
            font-size: 10px;
            font-weight: 800;
            white-space: nowrap;
        }}

        .status.passed {{
            color: var(--green);
            background: #e6f6ee;
        }}

        .status.failed {{
            color: var(--red);
            background: #fdeaea;
        }}

        .status.other {{
            color: var(--blue);
            background: #eaf1fb;
        }}

        .empty {{
            padding: 35px;
            color: var(--muted);
            text-align: center;
        }}

        .footer {{
            margin-top: 22px;
            color: var(--muted);
            font-size: 12px;
            text-align: center;
        }}

        @media (max-width: 1000px) {{
            .summary {{
                grid-template-columns: repeat(3, minmax(0, 1fr));
            }}
        }}

        @media (max-width: 650px) {{
            body {{
                padding: 16px;
            }}

            .header {{
                flex-direction: column;
            }}

            h1 {{
                font-size: 25px;
            }}

            .summary {{
                grid-template-columns: repeat(2, minmax(0, 1fr));
                gap: 10px;
            }}

            .card {{
                padding: 15px;
            }}

            .card-value {{
                font-size: 22px;
            }}

            .panel-header {{
                padding: 17px;
            }}

            th, td {{
                padding: 12px;
            }}
        }}
    </style>
</head>

<body>
    <main class="container">
        <header class="header">
            <div>
                <div class="eyebrow">JE BikeZ · Quality Assurance</div>
                <h1>Automation Test Report</h1>
                <p class="subtitle">
                    Generated on {safe(generated_at)}
                </p>
            </div>

            <div>
                <span class="overall {overall_class}">
                    {safe(overall_status)}
                </span>
            </div>
        </header>

        <section class="summary">
            <div class="card">
                <div class="card-label">Total Tests</div>
                <div class="card-value">{total_tests}</div>
            </div>

            <div class="card">
                <div class="card-label">Passed</div>
                <div class="card-value green">{passed_tests}</div>
            </div>

            <div class="card">
                <div class="card-label">Failed</div>
                <div class="card-value red">{failed_tests}</div>
            </div>

            <div class="card">
                <div class="card-label">Other / Skipped</div>
                <div class="card-value">{other_tests}</div>
            </div>

            <div class="card">
                <div class="card-label">Duration</div>
                <div class="card-value">
                    {safe(format_duration(total_duration))}
                </div>
            </div>
        </section>

        <section class="panel">
            <div class="panel-header">
                <div>
                    <h2>Test Results</h2>
                    <div class="panel-subtitle">
                        Overall results for each executed test
                    </div>
                </div>
            </div>

            <div class="table-wrap">
                <table>
                    <thead>
                        <tr>
                            <th>Test Name</th>
                            <th>Status</th>
                            <th>Duration</th>
                            <th>Message</th>
                        </tr>
                    </thead>
                    <tbody>
                        {result_rows}
                    </tbody>
                </table>
            </div>
        </section>

        <footer class="footer">
            JE BikeZ Automation Framework · HTML Execution Report
        </footer>
    </main>
</body>
</html>
"""

    output_path.write_text(html_document, encoding="utf-8")
    logger.info("HTML report saved: %s", output_path)

    return str(output_path)
